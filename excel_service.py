"""
Capa de acceso al Excel. Todo lo que la app necesita leer o escribir
sobre Sistema_Etoile.xlsx pasa por aquí.

Diseño:
- INVENTARIO es el catálogo maestro: CODIGO, PRENDA, COLOR, TALLA, PRECIO, STOCK
- HISTORIAL es el registro de cada prenda vendida (una fila por prenda, agrupadas
  por PEDIDO_ID cuando vienen del mismo pedido).
- El "disponible" de una prenda = STOCK (en INVENTARIO) - suma de CANT vendidas
  en HISTORIAL para ese mismo CODIGO. Así el stock en INVENTARIO nunca se
  edita automáticamente al vender; solo se resta en la lectura. Esto evita
  perder el número real de piezas que compraste si algo falla a mitad de un guardado.
"""
import os
import shutil
import threading
from datetime import datetime

import openpyxl

BASE_DIR = os.path.dirname(os.path.abspath(__file__))
EXCEL_PATH = os.environ.get(
    "ETOILE_EXCEL_PATH", os.path.join(BASE_DIR, "data", "Sistema_Etoile_2.xlsx")
)
BACKUP_DIR = os.path.join(BASE_DIR, "data", "backups")

_lock = threading.Lock()

INV_HEADER_ROW = 3
INV_FIRST_DATA_ROW = 4
INV_COLS = {"CODIGO": 1, "PRENDA": 2, "COLOR": 3, "TALLA": 4, "PRECIO": 5, "STOCK": 6}

HIST_HEADER_ROW = 4
HIST_FIRST_DATA_ROW = 5
# La hoja HISTORIAL tiene, más abajo, una sección de "RESUMEN" con fórmulas
# (empieza alrededor de la fila 211) que NO son pedidos reales. Los datos de
# pedidos reales viven solo entre HIST_FIRST_DATA_ROW y HIST_LAST_DATA_ROW,
# el mismo rango que usan las fórmulas de resumen del propio Excel
# (ej. SUMIFS($F$5:$F$204,...)). Si algún día se llenan más de estas filas,
# hay que ampliar este número (y las fórmulas del RESUMEN) antes de que choquen.
HIST_LAST_DATA_ROW = 204
HIST_COLS = {
    "FECHA": 1,
    "CODIGO": 2,
    "PRENDA": 3,
    "COLOR": 4,
    "TALLA": 5,
    "CANT": 6,
    "PRECIO": 7,
    "SUBTOTAL": 8,
    "CLIENTE": 9,
    "PEDIDO_ID": 10,
    "TELEFONO": 11,
    "METODO_PAGO": 12,
    "NUM_TRANSACCION": 13,
}


def _backup():
    os.makedirs(BACKUP_DIR, exist_ok=True)
    stamp = datetime.now().strftime("%Y%m%d_%H%M%S")
    dest = os.path.join(BACKUP_DIR, f"Sistema_Etoile_{stamp}.xlsx")
    try:
        shutil.copy2(EXCEL_PATH, dest)
        # mantener solo los últimos 30 backups
        backups = sorted(
            f for f in os.listdir(BACKUP_DIR) if f.startswith("Sistema_Etoile_")
        )
        for old in backups[:-30]:
            os.remove(os.path.join(BACKUP_DIR, old))
    except FileNotFoundError:
        pass


def _load(data_only=False):
    return openpyxl.load_workbook(EXCEL_PATH, data_only=data_only)


import re

_REF_SIMPLE = re.compile(r"^=([A-Z]+)(\d+)$")


def _resolver_celda(ws, row, col, visitados=None):
    """
    Si la celda tiene una fórmula simple de referencia directa (ej. '=B29'),
    sigue la cadena y devuelve el valor literal. Si es una fórmula más
    compleja (que no debería aparecer en INVENTARIO), devuelve None en vez
    de romper la app.
    """
    from openpyxl.utils import column_index_from_string

    if visitados is None:
        visitados = set()
    val = ws.cell(row=row, column=col).value
    if not (isinstance(val, str) and val.startswith("=")):
        return val
    m = _REF_SIMPLE.match(val)
    if not m:
        return None
    ref_col = column_index_from_string(m.group(1))
    ref_row = int(m.group(2))
    if (ref_row, ref_col) in visitados:
        return None
    visitados.add((ref_row, ref_col))
    return _resolver_celda(ws, ref_row, ref_col, visitados)


def get_inventario_raw(wb=None):
    """Lee INVENTARIO tal cual está en el Excel (sin calcular disponible)."""
    own = wb is None
    wb = wb or _load()
    ws = wb["INVENTARIO"]
    items = []
    for row in range(INV_FIRST_DATA_ROW, ws.max_row + 1):
        codigo = ws.cell(row=row, column=INV_COLS["CODIGO"]).value
        if not codigo:
            continue
        items.append(
            {
                "codigo": str(codigo).strip(),
                "prenda": _resolver_celda(ws, row, INV_COLS["PRENDA"]),
                "color": _resolver_celda(ws, row, INV_COLS["COLOR"]),
                "talla": _resolver_celda(ws, row, INV_COLS["TALLA"]),
                "precio": ws.cell(row=row, column=INV_COLS["PRECIO"]).value or 0,
                "stock": ws.cell(row=row, column=INV_COLS["STOCK"]).value or 0,
                "_row": row,
            }
        )
    if own:
        wb.close()
    return items


def get_historial_raw(wb=None):
    own = wb is None
    wb = wb or _load()
    ws = wb["HISTORIAL"]
    items = []
    for row in range(HIST_FIRST_DATA_ROW, HIST_LAST_DATA_ROW + 1):
        codigo = ws.cell(row=row, column=HIST_COLS["CODIGO"]).value
        if not codigo or (isinstance(codigo, str) and codigo.startswith("=")):
            continue
        items.append(
            {
                "fecha": ws.cell(row=row, column=HIST_COLS["FECHA"]).value,
                "codigo": str(codigo).strip(),
                "prenda": ws.cell(row=row, column=HIST_COLS["PRENDA"]).value,
                "color": ws.cell(row=row, column=HIST_COLS["COLOR"]).value,
                "talla": ws.cell(row=row, column=HIST_COLS["TALLA"]).value,
                "cant": ws.cell(row=row, column=HIST_COLS["CANT"]).value or 0,
                "precio": ws.cell(row=row, column=HIST_COLS["PRECIO"]).value or 0,
                "subtotal": ws.cell(row=row, column=HIST_COLS["SUBTOTAL"]).value or 0,
                "cliente": ws.cell(row=row, column=HIST_COLS["CLIENTE"]).value,
                "pedido_id": ws.cell(row=row, column=HIST_COLS["PEDIDO_ID"]).value,
                "telefono": ws.cell(row=row, column=HIST_COLS["TELEFONO"]).value,
                "metodo_pago": ws.cell(row=row, column=HIST_COLS["METODO_PAGO"]).value,
                "num_transaccion": ws.cell(
                    row=row, column=HIST_COLS["NUM_TRANSACCION"]
                ).value,
            }
        )
    if own:
        wb.close()
    return items


def get_inventario_con_disponible():
    """Devuelve INVENTARIO con 'disponible' calculado = stock - vendido."""
    with _lock:
        wb = _load()
        inventario = get_inventario_raw(wb)
        historial = get_historial_raw(wb)
        wb.close()

    vendido_por_codigo = {}
    for h in historial:
        vendido_por_codigo[h["codigo"]] = vendido_por_codigo.get(h["codigo"], 0) + (
            h["cant"] or 0
        )

    for item in inventario:
        vendido = vendido_por_codigo.get(item["codigo"], 0)
        item["vendido"] = vendido
        item["disponible"] = max(0, (item["stock"] or 0) - vendido)

    return inventario


def get_prendas_agrupadas(busqueda=None):
    """
    Agrupa el inventario por PRENDA -> lista de variantes (color, talla, precio, disponible).
    Si se pasa busqueda, filtra por nombre de prenda (contiene, insensible a mayúsculas).
    """
    inventario = get_inventario_con_disponible()
    if busqueda:
        b = busqueda.strip().lower()
        inventario = [i for i in inventario if b in (i["prenda"] or "").lower()]

    agrupado = {}
    for item in inventario:
        nombre = item["prenda"] or "(sin nombre)"
        agrupado.setdefault(nombre, []).append(item)

    resultado = []
    for nombre, variantes in sorted(agrupado.items()):
        colores = sorted(set(v["color"] for v in variantes if v["color"]))
        tallas = sorted(set(v["talla"] for v in variantes if v["talla"]))
        total_disponible = sum(v["disponible"] for v in variantes)
        precio = variantes[0]["precio"] if variantes else 0
        resultado.append(
            {
                "prenda": nombre,
                "colores": colores,
                "tallas": tallas,
                "precio": precio,
                "total_disponible": total_disponible,
                "variantes": variantes,
            }
        )
    return resultado


def actualizar_stock(codigo, nuevo_stock):
    with _lock:
        wb = _load()
        ws = wb["INVENTARIO"]
        encontrado = False
        for row in range(INV_FIRST_DATA_ROW, ws.max_row + 1):
            c = ws.cell(row=row, column=INV_COLS["CODIGO"]).value
            if c and str(c).strip() == str(codigo).strip():
                ws.cell(row=row, column=INV_COLS["STOCK"], value=int(nuevo_stock))
                encontrado = True
                break
        if not encontrado:
            wb.close()
            raise ValueError(f"Código {codigo} no encontrado en INVENTARIO")
        _backup()
        wb.save(EXCEL_PATH)
        wb.close()
    return True


def _abreviar(texto, largo):
    """Toma las primeras letras alfanuméricas de un texto y las pone en mayúsculas."""
    limpio = "".join(ch for ch in (texto or "") if ch.isalnum())
    return limpio[:largo].upper() if limpio else "XX"


def sugerir_codigo(prenda, color, talla, wb=None):
    """
    Genera un código sugerido siguiendo el mismo patrón que ya usa el Excel,
    ej. 'Blusa Eliana' + 'Vino' + 'S' -> 'BLEL-VI-S'.
    Si ya existe, le agrega un número al final hasta encontrar uno libre.
    """
    palabras = [p for p in (prenda or "").split() if p]
    if len(palabras) >= 2:
        prendaAbrev = _abreviar(palabras[0], 2) + _abreviar(palabras[1], 2)
    elif len(palabras) == 1:
        prendaAbrev = _abreviar(palabras[0], 4)
    else:
        prendaAbrev = "XXXX"

    colorAbrev = _abreviar(color, 2)
    tallaTxt = "".join(ch for ch in (talla or "").strip().upper() if ch.isalnum()) or "XX"

    base = f"{prendaAbrev}-{colorAbrev}-{tallaTxt}"

    own = wb is None
    wb = wb or _load()
    inventario = get_inventario_raw(wb)
    if own:
        wb.close()
    existentes = {i["codigo"] for i in inventario}

    if base not in existentes:
        return base

    n = 2
    while f"{base}-{n}" in existentes:
        n += 1
    return f"{base}-{n}"


def get_prendas_nombres():
    """Nombres de prenda ya existentes, para sugerir/reusar en el formulario de alta."""
    inventario = get_inventario_raw()
    nombres = sorted(
        set(
            i["prenda"]
            for i in inventario
            if i["prenda"] and not str(i["prenda"]).startswith("=")
        )
    )
    return nombres


def agregar_prenda(prenda, color, talla, precio, stock, codigo=None):
    """
    Agrega una nueva variante (fila) a INVENTARIO. Si no se da codigo, se
    autogenera. Si el codigo dado ya existe, se rechaza (para no pisar una
    variante existente sin querer; para eso está actualizar_stock).
    """
    prenda = (prenda or "").strip()
    color = (color or "").strip()
    talla = (talla or "").strip()

    if not prenda:
        raise ValueError("Falta el nombre de la prenda")
    if precio is None or float(precio) < 0:
        raise ValueError("Precio inválido")
    if stock is None or int(stock) < 0:
        raise ValueError("Stock inválido")

    with _lock:
        wb = _load()
        ws = wb["INVENTARIO"]
        inventario = get_inventario_raw(wb)
        existentes = {i["codigo"] for i in inventario}

        if codigo:
            codigo = codigo.strip().upper()
            if codigo in existentes:
                wb.close()
                raise ValueError(
                    f"El código {codigo} ya existe. Usa 'Actualizar stock' si "
                    "quieres modificar esa variante, o elige otro código."
                )
        else:
            codigo = sugerir_codigo(prenda, color, talla, wb)

        # primera fila libre después de los datos existentes
        fila = INV_FIRST_DATA_ROW
        for row in range(INV_FIRST_DATA_ROW, ws.max_row + 2):
            if not ws.cell(row=row, column=INV_COLS["CODIGO"]).value:
                fila = row
                break
        else:
            fila = ws.max_row + 1

        ws.cell(row=fila, column=INV_COLS["CODIGO"], value=codigo)
        ws.cell(row=fila, column=INV_COLS["PRENDA"], value=prenda)
        ws.cell(row=fila, column=INV_COLS["COLOR"], value=color)
        ws.cell(row=fila, column=INV_COLS["TALLA"], value=talla)
        ws.cell(row=fila, column=INV_COLS["PRECIO"], value=float(precio))
        ws.cell(row=fila, column=INV_COLS["STOCK"], value=int(stock))

        _backup()
        wb.save(EXCEL_PATH)
        wb.close()

    return codigo


def siguiente_pedido_id(wb):
    ws = wb["HISTORIAL"]
    max_id = 0
    for row in range(HIST_FIRST_DATA_ROW, HIST_LAST_DATA_ROW + 1):
        pid = ws.cell(row=row, column=HIST_COLS["PEDIDO_ID"]).value
        if pid:
            try:
                num = int(str(pid).replace("PED-", ""))
                max_id = max(max_id, num)
            except ValueError:
                continue
    return f"PED-{max_id + 1:04d}"


def crear_pedido(cliente, telefono, items, metodo_pago, num_transaccion):
    """
    items: lista de {codigo, prenda, color, talla, precio, cantidad}
    metodo_pago: 'efectivo' | 'transferencia'
    num_transaccion: string o None (obligatorio si es transferencia)
    Valida disponibilidad antes de escribir. Devuelve el pedido_id creado.
    """
    if not items:
        raise ValueError("El pedido no tiene prendas")
    if metodo_pago not in ("efectivo", "transferencia"):
        raise ValueError("Método de pago inválido")
    if metodo_pago == "transferencia" and not num_transaccion:
        raise ValueError("Falta el número de transacción")

    with _lock:
        wb = _load()
        ws_hist = wb["HISTORIAL"]
        inventario = get_inventario_raw(wb)
        historial = get_historial_raw(wb)

        stock_por_codigo = {i["codigo"]: i["stock"] for i in inventario}
        vendido_por_codigo = {}
        for h in historial:
            vendido_por_codigo[h["codigo"]] = vendido_por_codigo.get(h["codigo"], 0) + (
                h["cant"] or 0
            )

        # validar disponibilidad
        for it in items:
            codigo = it["codigo"]
            disponible = stock_por_codigo.get(codigo, 0) - vendido_por_codigo.get(codigo, 0)
            if it["cantidad"] > disponible:
                wb.close()
                raise ValueError(
                    f"Stock insuficiente para {codigo} ({it.get('prenda','')} "
                    f"{it.get('color','')} {it.get('talla','')}): "
                    f"disponible {disponible}, pedido {it['cantidad']}"
                )

        pedido_id = siguiente_pedido_id(wb)
        fecha = datetime.now().strftime("%Y-%m-%d %H:%M")

        # buscar la primera fila realmente libre dentro del rango de datos reales
        fila = None
        for row in range(HIST_FIRST_DATA_ROW, HIST_LAST_DATA_ROW + 1):
            valor = ws_hist.cell(row=row, column=HIST_COLS["CODIGO"]).value
            if not valor or (isinstance(valor, str) and valor.startswith("=")):
                fila = row
                break

        filas_libres = fila is not None and (
            HIST_LAST_DATA_ROW - fila + 1
        ) >= len(items)

        if fila is None or not filas_libres:
            wb.close()
            raise ValueError(
                "La hoja HISTORIAL se quedó sin filas libres en su rango de datos "
                f"(hasta la fila {HIST_LAST_DATA_ROW}). Hay que ampliar el rango "
                "en el Excel (incluyendo las fórmulas de RESUMEN) antes de "
                "seguir registrando pedidos."
            )

        for it in items:
            subtotal = round(it["precio"] * it["cantidad"], 2)
            ws_hist.cell(row=fila, column=HIST_COLS["FECHA"], value=fecha)
            ws_hist.cell(row=fila, column=HIST_COLS["CODIGO"], value=it["codigo"])
            ws_hist.cell(row=fila, column=HIST_COLS["PRENDA"], value=it.get("prenda"))
            ws_hist.cell(row=fila, column=HIST_COLS["COLOR"], value=it.get("color"))
            ws_hist.cell(row=fila, column=HIST_COLS["TALLA"], value=it.get("talla"))
            ws_hist.cell(row=fila, column=HIST_COLS["CANT"], value=it["cantidad"])
            ws_hist.cell(row=fila, column=HIST_COLS["PRECIO"], value=it["precio"])
            ws_hist.cell(row=fila, column=HIST_COLS["SUBTOTAL"], value=subtotal)
            ws_hist.cell(row=fila, column=HIST_COLS["CLIENTE"], value=cliente)
            ws_hist.cell(row=fila, column=HIST_COLS["PEDIDO_ID"], value=pedido_id)
            ws_hist.cell(row=fila, column=HIST_COLS["TELEFONO"], value=telefono)
            ws_hist.cell(row=fila, column=HIST_COLS["METODO_PAGO"], value=metodo_pago)
            ws_hist.cell(
                row=fila,
                column=HIST_COLS["NUM_TRANSACCION"],
                value=num_transaccion if metodo_pago == "transferencia" else "EFECTIVO",
            )
            fila += 1

        _backup()
        wb.save(EXCEL_PATH)
        wb.close()

    return pedido_id


def get_pedidos_agrupados(desde=None, hasta=None, cliente=None, codigo=None):
    """Agrupa HISTORIAL por pedido_id para mostrar en la vista de historial."""
    historial = get_historial_raw()

    def pasa_filtros(h):
        if desde and (not h["fecha"] or str(h["fecha"])[:10] < desde):
            return False
        if hasta and (not h["fecha"] or str(h["fecha"])[:10] > hasta):
            return False
        if cliente and cliente.lower() not in (h["cliente"] or "").lower():
            return False
        if codigo and codigo.lower() not in (h["codigo"] or "").lower():
            return False
        return True

    historial = [h for h in historial if pasa_filtros(h)]

    pedidos = {}
    for h in historial:
        pid = h["pedido_id"] or f"SIN-ID-{h['fecha']}-{h['cliente']}"
        if pid not in pedidos:
            pedidos[pid] = {
                "pedido_id": pid,
                "fecha": h["fecha"],
                "cliente": h["cliente"],
                "telefono": h["telefono"],
                "metodo_pago": h["metodo_pago"],
                "num_transaccion": h["num_transaccion"],
                "items": [],
                "total": 0,
            }
        pedidos[pid]["items"].append(h)
        pedidos[pid]["total"] += h["subtotal"] or 0

    resultado = list(pedidos.values())
    resultado.sort(key=lambda p: p["fecha"] or "", reverse=True)
    return resultado
